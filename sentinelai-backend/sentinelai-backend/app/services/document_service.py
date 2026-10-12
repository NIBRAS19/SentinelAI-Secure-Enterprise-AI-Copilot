"""
Document upload + ingestion pipeline.

State machine: UPLOADED -> VALIDATING -> EXTRACTING -> CHUNKING -> EMBEDDING
-> INDEXING -> READY, or FAILED at any step (with error_message set and
retry_count incremented).

Quarantine: if scan_for_prompt_injection() flags the extracted text as
SUSPICIOUS, the document's chunks are still written (so they're visible in
the UI for an admin to review) but embedding is skipped entirely — chunks
are persisted with embedding=NULL. Since retrieval always filters on
`Chunk.embedding IS NOT NULL`, a SUSPICIOUS document structurally cannot
surface in RAG results. This is enforced by what gets written to the
database, not by a status check retrieval code has to remember to apply.
"""
import hashlib
from datetime import datetime, timezone

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.ai.embeddings import get_embedding_provider
from app.ai.prompt_injection import scan_for_prompt_injection
from app.ai.text import chunk_text, normalize_whitespace
from app.core.errors import AppError, NotFoundError
from app.models.chunk import Chunk
from app.models.document import Document
from app.services import audit_service, storage

_UNSET = object()


def _extract_text(mime_type: str, data: bytes) -> tuple[str, int | None]:
    """Returns (plain_text, page_count or None)."""
    if mime_type == "application/pdf":
        from pypdf import PdfReader
        import io

        reader = PdfReader(io.BytesIO(data))
        pages_text = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages_text), len(reader.pages)

    if mime_type in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ):
        import docx
        import io

        doc = docx.Document(io.BytesIO(data))
        return "\n\n".join(p.text for p in doc.paragraphs), None

    # Plain text / markdown fallback.
    return data.decode("utf-8", errors="replace"), None


async def create_document(
    db: AsyncSession,
    *,
    organization_id: str,
    uploaded_by_id: str,
    collection_id: str | None,
    name: str,
    mime_type: str,
    data: bytes,
) -> Document:
    checksum = hashlib.sha256(data).hexdigest()
    document = Document(
        organization_id=organization_id,
        collection_id=collection_id,
        uploaded_by_id=uploaded_by_id,
        name=name,
        mime_type=mime_type,
        size_bytes=len(data),
        storage_key="",  # filled in below once we have the document id
        checksum_sha256=checksum,
        status="UPLOADED",
        security_status="SAFE",
    )
    db.add(document)
    await db.flush()

    storage_key = storage.build_storage_key(organization_id, document.id, name)
    storage.save_bytes(storage_key, data)
    document.storage_key = storage_key
    await db.flush()

    await audit_service.record_event(
        db,
        organization_id=organization_id,
        actor_id=uploaded_by_id,
        actor_label=uploaded_by_id,
        action="document.upload",
        resource_type="document",
        resource_id=document.id,
        details={"name": name, "size_bytes": len(data)},
    )
    return document


async def process_document(db: AsyncSession, *, document_id: str, organization_id: str) -> None:
    document = (
        await db.execute(select(Document).where(Document.id == document_id, Document.organization_id == organization_id))
    ).scalar_one_or_none()
    if document is None:
        return

    document.last_attempt_at = datetime.now(timezone.utc)

    try:
        document.status = "VALIDATING"
        await db.flush()
        raw_bytes = storage.read_bytes(document.storage_key)

        document.status = "EXTRACTING"
        await db.flush()
        raw_text, page_count = _extract_text(document.mime_type, raw_bytes)
        normalized = normalize_whitespace(raw_text)
        document.page_count = page_count

        injection = scan_for_prompt_injection(normalized)
        if injection.is_suspicious:
            document.security_status = "SUSPICIOUS"
            await audit_service.record_event(
                db,
                organization_id=organization_id,
                actor_id=None,
                actor_label="system:ingestion",
                action="document.quarantined",
                resource_type="document",
                resource_id=document.id,
                severity="WARNING",
                details={
                    "matched_patterns": injection.matched_patterns,
                    "flagged_excerpt": _flagged_snippet(normalized, injection.matched_snippets[0] if injection.matched_snippets else ""),
                },
            )

        document.status = "CHUNKING"
        await db.flush()
        pieces = chunk_text(normalized)

        document.status = "EMBEDDING" if document.security_status == "SAFE" else "INDEXING"
        await db.flush()

        provider = get_embedding_provider()
        embeddings: list[list[float] | None]
        if document.security_status == "SAFE":
            embeddings = await provider.embed_batch(pieces)
        else:
            # Quarantined: persist chunks for visibility, but never embed them.
            embeddings = [None] * len(pieces)

        for index, (content, embedding) in enumerate(zip(pieces, embeddings)):
            db.add(
                Chunk(
                    organization_id=organization_id,
                    document_id=document.id,
                    content=content,
                    chunk_index=index,
                    page_number=None,
                    embedding=embedding,
                )
            )

        document.chunk_count = len(pieces)
        document.status = "INDEXING"
        await db.flush()

        document.status = "READY"
        document.error_message = None
        await db.flush()

    except Exception as exc:  # noqa: BLE001 - ingestion must never crash the worker; record and move on
        document.status = "FAILED"
        document.error_message = str(exc)[:2000]
        document.retry_count += 1
        await db.flush()
        await audit_service.record_event(
            db,
            organization_id=organization_id,
            actor_id=None,
            actor_label="system:ingestion",
            action="document.ingestion_failed",
            resource_type="document",
            resource_id=document.id,
            severity="WARNING",
            details={"error": str(exc)[:500]},
        )


def _flagged_snippet(text: str, matched_snippet: str, window: int = 160) -> str:
    """
    Builds a short excerpt around the ACTUAL matched text (not the regex
    pattern source — a regex pattern string like
    "disregard (all|any|previous|prior|the above)" is never a literal
    substring of the document, so searching for it would never find
    anything). `matched_snippet` must come from InjectionScanResult's
    matched_snippets (the real `m.group(0)` text), not matched_patterns.
    """
    position = text.lower().find(matched_snippet.lower()) if matched_snippet else -1
    if position == -1:
        return text[:400].strip()
    start = max(0, position - window)
    end = min(len(text), position + len(matched_snippet) + window)
    prefix = "..." if start > 0 else ""
    suffix = "..." if end < len(text) else ""
    return f"{prefix}{text[start:end].strip()}{suffix}"


async def get_document(db: AsyncSession, *, organization_id: str, document_id: str) -> Document:
    document = (
        await db.execute(select(Document).where(Document.id == document_id, Document.organization_id == organization_id))
    ).scalar_one_or_none()
    if document is None:
        raise NotFoundError("Document not found.")
    return document


async def list_documents(db: AsyncSession, *, organization_id: str, collection_id: str | None = None) -> list[Document]:
    stmt = select(Document).where(Document.organization_id == organization_id)
    if collection_id:
        stmt = stmt.where(Document.collection_id == collection_id)
    stmt = stmt.order_by(Document.created_at.desc())
    return list((await db.execute(stmt)).scalars().all())


async def update_document(
    db: AsyncSession,
    *,
    organization_id: str,
    document_id: str,
    collection_id: str | None = _UNSET,
    access_roles: list[str] | None = _UNSET,
) -> Document:
    document = await get_document(db, organization_id=organization_id, document_id=document_id)
    if collection_id is not _UNSET:
        document.collection_id = collection_id
    if access_roles is not _UNSET:
        document.access_roles = access_roles
    await db.flush()
    return document


async def delete_document(db: AsyncSession, *, organization_id: str, document_id: str) -> None:
    document = await get_document(db, organization_id=organization_id, document_id=document_id)
    try:
        storage.delete(document.storage_key)
    except Exception:  # noqa: BLE001 - best-effort file cleanup; row deletion still proceeds
        pass
    await db.delete(document)
    await db.flush()
